# services/export_service.py
#
# Builds downloadable exports of the Literature Matrix (see routes/papers_routes.py's
# export_matrix_excel() / export_matrix_pdf() / export_matrix_docx(), and
# templates/literature_matrix.html's "Export" dropdown) - an Excel workbook via
# openpyxl, a PDF via reportlab, and a Word document via python-docx, all built from
# the same saved-papers data so the three formats never drift apart.
#
# Every builder returns an in-memory BytesIO, never a temp file on disk, so they drop
# straight into Flask's send_file().

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
from xml.sax.saxutils import escape as _xml_escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

COLUMNS = ["Paper", "Authors", "Year", "Methodology", "Sample", "Findings", "Limitations"]

NAVY = "1E3A8A"
NOT_EXTRACTED = "Not extracted yet"


def _row_for_paper(paper):
    """
    One matrix row's plain-text values, in COLUMNS order - the single source of truth
    both build_matrix_excel() and build_matrix_pdf() read from, so a paper missing a
    field renders the same placeholder text in both exports as it does on-page in
    literature_matrix.html.
    """
    return [
        paper.title,
        paper.authors or "Unknown authors",
        paper.year or "",
        paper.matrix_methodology or NOT_EXTRACTED,
        paper.matrix_sample or NOT_EXTRACTED,
        paper.matrix_findings or NOT_EXTRACTED,
        paper.matrix_limitations or NOT_EXTRACTED,
    ]


def _slugify(text):
    """Turn a project title into a safe download filename fragment."""
    slug = re.sub(r"[^A-Za-z0-9]+", "_", text or "project").strip("_")
    return slug or "project"


def export_filename(project, extension):
    """e.g. 'AI_in_Academic_Research_Workflows_Literature_Matrix.xlsx'."""
    return f"{_slugify(project.title)}_Literature_Matrix.{extension}"


def build_matrix_excel(project, papers) -> io.BytesIO:
    """
    An .xlsx workbook of the Literature Matrix: a title/meta header, then one row per
    saved paper with columns matching literature_matrix.html's table.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Literature Matrix"

    ws.append([f"Literature Matrix – {project.title}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    ws.cell(row=1, column=1).font = Font(size=14, bold=True)

    paper_word = "paper" if len(papers) == 1 else "papers"
    ws.append([f"Exported {datetime.utcnow().strftime('%B %d, %Y')} · {len(papers)} {paper_word}"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNS))
    ws.cell(row=2, column=1).font = Font(size=10, italic=True, color="6B7280")

    ws.append([])  # blank spacer row

    header_row = 4
    ws.append(COLUMNS)
    header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    for col in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for paper in papers:
        ws.append(_row_for_paper(paper))

    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, min_col=1, max_col=len(COLUMNS)):
        for cell in row:
            cell.alignment = wrap

    for i, width in enumerate([32, 20, 8, 34, 30, 34, 30], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[header_row].height = 22
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _rl_text(value):
    """
    reportlab's Paragraph parses its text as a small XML/HTML-like markup language, so
    raw AI-extracted or user-edited text (which can legitimately contain '&', '<', '>')
    would otherwise throw a parse error or render wrong. Escape first, then reinstate
    line breaks as the one tag we intentionally allow through.
    """
    text = _xml_escape(str(value if value not in (None, "") else ""))
    return text.replace("\n", "<br/>")


def build_matrix_pdf(project, papers) -> io.BytesIO:
    """
    A landscape PDF of the same Literature Matrix data as build_matrix_excel(). Cells
    are reportlab Paragraphs (not bare strings) so long extracted text wraps within its
    column instead of overflowing the page.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        leftMargin=0.4 * inch, rightMargin=0.4 * inch,
        topMargin=0.5 * inch, bottomMargin=0.5 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("MatrixTitle", parent=styles["Heading1"], fontSize=16, spaceAfter=2)
    meta_style = ParagraphStyle(
        "MatrixMeta", parent=styles["Normal"], fontSize=9,
        textColor=colors.HexColor("#6B7280"), spaceAfter=14,
    )
    header_style = ParagraphStyle(
        "MatrixHeader", parent=styles["Normal"], fontSize=9,
        textColor=colors.white, fontName="Helvetica-Bold",
    )
    cell_style = ParagraphStyle("MatrixCell", parent=styles["Normal"], fontSize=8.5, leading=11)

    paper_word = "paper" if len(papers) == 1 else "papers"
    elements = [
        Paragraph(f"Literature Matrix &ndash; {_xml_escape(project.title)}", title_style),
        Paragraph(
            f"Exported {datetime.utcnow().strftime('%B %d, %Y')} &middot; {len(papers)} {paper_word}",
            meta_style,
        ),
    ]

    header = [Paragraph(col, header_style) for col in COLUMNS]
    data = [header]
    for paper in papers:
        data.append([Paragraph(_rl_text(value), cell_style) for value in _row_for_paper(paper)])

    col_widths = [1.5 * inch, 1.1 * inch, 0.45 * inch, 1.7 * inch, 1.5 * inch, 1.7 * inch, 1.5 * inch]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def _shade_cell(cell, hex_color):
    """python-docx has no cell-background API - set it via the cell's own XML directly."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _mark_repeat_header(row):
    """Flags a table row to repeat as a header on every page it spans across."""
    row_props = row._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    row_props.append(header_flag)


DOCX_COL_WIDTHS = [Inches(1.7), Inches(1.3), Inches(0.6), Inches(1.9), Inches(1.7), Inches(1.9), Inches(1.7)]


def build_matrix_docx(project, papers) -> io.BytesIO:
    """
    A landscape Word document of the same Literature Matrix data as build_matrix_excel()
    and build_matrix_pdf(): a title, a meta line, then a 7-column table with a shaded,
    repeating header row. No markup-escaping is needed here (unlike the PDF) - python-docx
    writes cell text as plain text runs, not through a markup mini-language.
    """
    doc = Document()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Inches(0.5)
    section.top_margin = section.bottom_margin = Inches(0.5)

    title = doc.add_heading(level=1)
    title_run = title.add_run(f"Literature Matrix – {project.title}")
    title_run.font.size = Pt(18)

    paper_word = "paper" if len(papers) == 1 else "papers"
    meta = doc.add_paragraph()
    meta_run = meta.add_run(f"Exported {datetime.utcnow().strftime('%B %d, %Y')} · {len(papers)} {paper_word}")
    meta_run.italic = True
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    doc.add_paragraph()  # spacer before the table

    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    header_cells = table.rows[0].cells
    for i, (col_name, width) in enumerate(zip(COLUMNS, DOCX_COL_WIDTHS)):
        header_cells[i].width = width
        header_cells[i].text = ""
        run = header_cells[i].paragraphs[0].add_run(col_name)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade_cell(header_cells[i], NAVY)
    _mark_repeat_header(table.rows[0])

    for paper in papers:
        row_cells = table.add_row().cells
        for i, (value, width) in enumerate(zip(_row_for_paper(paper), DOCX_COL_WIDTHS)):
            row_cells[i].width = width
            row_cells[i].text = str(value if value not in (None, "") else "")
            for run in row_cells[i].paragraphs[0].runs:
                run.font.size = Pt(9.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
